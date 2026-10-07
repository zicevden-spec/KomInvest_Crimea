# -*- coding: utf-8 -*-
import os, json, io, traceback
from datetime import datetime
from telegram import (Update, ReplyKeyboardMarkup, KeyboardButton,
                      InlineKeyboardButton, InlineKeyboardMarkup, InputMediaPhoto)
from telegram.ext import (Application, CommandHandler, MessageHandler,
                          CallbackQueryHandler, ContextTypes, filters)
from dotenv import load_dotenv
import gh_store
import catalog

load_dotenv()

BOT_TOKEN = os.getenv("TELEGRAM_BOT_TOKEN")
ADMIN_ID = int(os.getenv("TELEGRAM_USER_ID", "0"))
CHANNEL_URL = os.getenv("CHANNEL_URL", "https://t.me/KomInvest_Crimea")
CHANNEL_ID = os.getenv("CHANNEL_ID", "@KomInvest_Crimea")
WEBHOOK_URL = os.getenv("WEBHOOK_URL") or os.getenv("RENDER_EXTERNAL_URL", "")
PORT = int(os.getenv("PORT", "8080"))
BOT_USERNAME = ""

OBJ_TYPES = ["офис", "склад", "стрит-ритейл", "торговый центр", "земельный участок", "другое"]
EDIT_FIELDS = {"title": "Заголовок", "desc": "Описание", "price": "Цена",
               "loc": "Локация", "area": "Площадь"}

user_states = {}

LEAD_STATUSES = {
    "new": "🔥 Новый",
    "work": "📞 В работе",
    "done": "✅ Закрыт",
    "fail": "❌ Отказ",
    "recall": "🕐 Не дозвон",
}

def update_leads(leads):
    gh_store.push("leads.json", leads)

def lead_card_text(l):
    return (f"👤 {l['name']}\n📞 {l['phone']}\n📍 Источник: {l['source']}\n"
            f"🕒 Принят: {l['date']}\n"
            f"📊 Статус: {LEAD_STATUSES.get(l['status'], l['status'])}"
            + (f"\n🔄 Изменён: {l['status_date']}" if l.get('status_date') else ""))

def lead_kb(l):
    lid = l["id"]
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📞 В работе", callback_data=f"ls_{lid}_work"),
         InlineKeyboardButton("✅ Закрыт", callback_data=f"ls_{lid}_done")],
        [InlineKeyboardButton("❌ Отказ", callback_data=f"ls_{lid}_fail"),
         InlineKeyboardButton("🕐 Не дозвон", callback_data=f"ls_{lid}_recall")],
        [InlineKeyboardButton("📱 Позвонить", url=f"tel:{l['phone']}")],
        [InlineKeyboardButton("⬅️ К лидам", callback_data="adm_leads")],
    ])

# ---------- Хранилище ----------

def load_leads():
    return gh_store.read_local("leads.json", [])

def save_lead(lead):
    leads = load_leads(); leads.append(lead)
    gh_store.push("leads.json", leads)

def load_admins():
    return gh_store.read_local("admins.json", {"super": [ADMIN_ID], "admins": []})

def save_admins(data):
    gh_store.push("admins.json", data)

def load_objects():
    return gh_store.read_local("objects.json", [])

def save_objects(objs):
    gh_store.push("objects.json", objs)

def is_super(uid):
    return uid in load_admins().get("super", [])

def is_admin_uid(uid):
    d = load_admins()
    return uid in d.get("super", []) or uid in d.get("admins", [])

# ---------- Клавиатуры ----------

def main_keyboard():
    return ReplyKeyboardMarkup(
        [[KeyboardButton("📝 Оставить заявку для юриста")],
         [KeyboardButton("ℹ️ О компании"), KeyboardButton("📰 Наш канал")]],
        resize_keyboard=True)

def keyboard_for(uid):
    rows = [
        [KeyboardButton("📝 Оставить заявку для юриста")],
        [KeyboardButton("ℹ️ О компании"), KeyboardButton("📰 Наш канал")]
    ]
    if is_admin_uid(uid):
        rows.append([KeyboardButton("🛠 Админка")])
    return ReplyKeyboardMarkup(rows, resize_keyboard=True)

# ---------- Клиент ----------

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    uid = update.effective_user.id
    user_states.pop(uid, None)
    context.user_data["source"] = context.args[0] if context.args else "direct"
    await update.message.reply_text(
        "Здравствуйте! 👋 Это бот «КомИнвест» — коммерческая недвижимость Крыма.\n\n"
        "Помогаем находить доходные объекты, проверять их юридическую чистоту "
        "и безопасно выходить на сделку.\n\n"
        "Нажмите кнопку ниже 👇",
        reply_markup=keyboard_for(uid))

async def about(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    await update.message.reply_text(
        "🏢 «КомИнвест» — подбор и проверка коммерческой недвижимости в Крыму.\n\n"
        "✅ Офисы, склады, стрит-ритейл, земельные участки\n"
        "✅ Юридическая проверка объекта перед покупкой\n"
        "✅ Сопровождение сделки «под ключ»\n\n"
        "Больше полезного — в нашем канале: " + CHANNEL_URL,
        reply_markup=keyboard_for(update.effective_user.id))

async def channel(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    await update.message.reply_text("📰 Подписывайтесь на наш канал:\n" + CHANNEL_URL,
                                    reply_markup=keyboard_for(update.effective_user.id))

async def lead_start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    user_states[update.effective_user.id] = "name"
    await update.message.reply_text("Отлично! 🙌\nКак к вам обращаться?",
                                    reply_markup=keyboard_for(update.effective_user.id))

async def finish_lead(update: Update, context: ContextTypes.DEFAULT_TYPE, phone: str):
    uid = update.effective_user.id
    name = context.user_data.get("name", "Без имени")
    lead = {"id": len(load_leads()) + 1, "name": name, "phone": phone,
            "source": context.user_data.get("source", "direct"),
            "date": datetime.now().strftime("%Y-%m-%d %H:%M"), "status": "новая"}
    save_lead(lead)
    user_states.pop(uid, None)
    context.user_data.pop("name", None)
    await update.message.reply_text(
        f"Спасибо, {name}! ✅\n\nВаша заявка принята. Юрист свяжется с вами "
        "в ближайшее время.\n\nА пока — загляните в наш канал:\n" + CHANNEL_URL,
        reply_markup=keyboard_for(uid))
    try:
        await context.bot.send_message(ADMIN_ID,
            f"🔥 НОВАЯ ЗАЯВКА #{lead['id']}\n\n👤 {name}\n📞 {phone}\n"
            f"📍 Источник: {lead['source']}\n\n/admin — админка | /export — Excel")
    except Exception as e:
        print(f"Не удалось уведомить админа: {e}")

async def handle_contact(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    await finish_lead(update, context, update.message.contact.phone_number)

# ---------- Объекты ----------

def obj_card(o):
    return (f"🏢 {o['title']}\n\n{o['description']}\n\n"
            f"💰 Цена: {o['price']}\n📍 Локация: {o['location']}\n"
            f"📐 Площадь: {o.get('area', '-')}\n🏷 Тип: {o.get('type', '-')}\n"
            f"🖼 Обложка: {'есть' if o.get('cover') else 'нет'}\n"
            f"{'✅ Опубликован' if o.get('published') else '⏳ Не опубликован'}")

def obj_view_kb(oid):
    return InlineKeyboardMarkup([
        [InlineKeyboardButton("📢 Опубликовать в канал", callback_data=f"obj_pub_{oid}")],
        [InlineKeyboardButton("✏️ Редактировать", callback_data=f"obj_edit_{oid}")],
        [InlineKeyboardButton("🗑 Удалить", callback_data=f"obj_del_{oid}")],
        [InlineKeyboardButton("⬅️ Список", callback_data="obj_list")],
    ])

def obj_edit_kb(oid):
    rows = [[InlineKeyboardButton(f"✏️ {label}", callback_data=f"oe_{f}_{oid}")]
            for f, label in EDIT_FIELDS.items()]
    rows.append([InlineKeyboardButton("🏷 Тип", callback_data=f"oe_type_{oid}")])
    rows.append([InlineKeyboardButton("🖼 Заменить обложку", callback_data=f"oe_cover_{oid}")])
    rows.append([InlineKeyboardButton("⬅️ К карточке", callback_data=f"obj_view_{oid}")])
    return InlineKeyboardMarkup(rows)

# ---------- Фото ----------

async def handle_photo(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    uid = update.effective_user.id
    st = user_states.get(uid)
    if not (isinstance(st, dict) and is_admin_uid(uid)):
        await update.message.reply_text("Сейчас не жду фото 😅", reply_markup=keyboard_for(uid))
        return
    step = st.get("step")
    fid = update.message.photo[-1].file_id

    if step == "obj_cover":
        st["cover"] = fid
        st["photos"] = []
        st["step"] = "obj_photos"
        await update.message.reply_text(
            "🖼 Обложка принята!\n\nТеперь пришлите дополнительные фото (до 9) — "
            "или напишите «далее».")
    elif step == "obj_photos":
        if len(st["photos"]) >= 9:
            await update.message.reply_text("Максимум 9 доп. фото. Напишите «далее».")
            return
        st["photos"].append(fid)
        await update.message.reply_text(f"📸 Доп. фото: {len(st['photos'])}/9. Ещё или «далее».")
    elif step == "obj_edit_cover":
        objs = load_objects()
        o = next((x for x in objs if x["id"] == st["oid"]), None)
        if o:
            o["cover"] = fid
            save_objects(objs)
        user_states.pop(uid, None)
        await update.message.reply_text("✅ Обложка заменена!", reply_markup=obj_view_kb(st["oid"]))
    else:
        await update.message.reply_text("Сейчас не жду фото 😅", reply_markup=keyboard_for(uid))

# ---------- Текст ----------

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    uid = update.effective_user.id
    text = update.message.text.strip()
    st = user_states.get(uid)
    step = st.get("step") if isinstance(st, dict) else st

    if isinstance(st, dict) and str(step).startswith("obj_") and is_admin_uid(uid):
        if step == "obj_photos":
            st["step"] = "obj_title"
            await update.message.reply_text("Теперь напишите ЗАГОЛОВОК объекта:")
            return
        if step == "obj_title":
            st["title"] = text; st["step"] = "obj_desc"
            await update.message.reply_text("📝 Описание объекта:")
            return
        if step == "obj_desc":
            st["description"] = text; st["step"] = "obj_price"
            await update.message.reply_text("💰 Цена (например: 12 500 000 ₽):")
            return
        if step == "obj_price":
            st["price"] = text; st["step"] = "obj_location"
            await update.message.reply_text("📍 Локация (город, район, улица):")
            return
        if step == "obj_location":
            st["location"] = text; st["step"] = "obj_area"
            await update.message.reply_text("📐 Площадь (например: 120 м²):")
            return
        if step == "obj_area":
            st["area"] = text; st["step"] = "obj_type"
            rows = [[InlineKeyboardButton(t, callback_data=f"ots_{i}_new")] for i, t in enumerate(OBJ_TYPES)]
            await update.message.reply_text("🏷 Выберите тип объекта:", reply_markup=InlineKeyboardMarkup(rows))
            return
        if step == "obj_edit":
            objs = load_objects()
            o = next((x for x in objs if x["id"] == st["oid"]), None)
            if o:
                key = {"title": "title", "desc": "description", "price": "price",
                       "loc": "location", "area": "area"}[st["field"]]
                o[key] = text
                save_objects(objs)
            user_states.pop(uid, None)
            await update.message.reply_text("✅ Сохранено! Карточка обновлена:",
                                            reply_markup=obj_view_kb(st["oid"]))
            if o:
                pic = o.get("cover") or (o.get("photos") or [None])[0]
                if pic:
                    try:
                        await update.message.reply_photo(photo=pic, caption=obj_card(o)[:1024])
                    except Exception as e:
                        await update.message.reply_text(obj_card(o))
                else:
                    await update.message.reply_text(obj_card(o))
            return

    if step in ("add_admin", "del_admin") and is_super(uid):
        try:
            target = (await context.bot.get_chat(text)).id if text.startswith("@") else int(text)
        except Exception:
            await update.message.reply_text("Не распознал. Пришли числовой ID или @username.",
                                            reply_markup=keyboard_for(uid))
            return
        d = load_admins()
        if step == "add_admin":
            if target not in d["admins"] and target not in d["super"]:
                d["admins"].append(target); save_admins(d); msg = f"✅ Админ добавлен: {target}"
            else: msg = "Этот человек уже админ."
        else:
            if target in d["admins"]:
                d["admins"].remove(target); save_admins(d); msg = f"🗑 Админ снят: {target}"
            else: msg = "Этого человека нет в админах."
        user_states.pop(uid, None)
        await update.message.reply_text(msg, reply_markup=admin_menu(uid))
        return

    if text in ("📝 Оставить заявку", "📝 Оставить заявку для юриста"):
        return await lead_start(update, context)
    if text == "ℹ️ О компании": return await about(update, context)
    if text == "📰 Наш канал": return await channel(update, context)
    if text == "🛠 Админка":
        return await admin_cmd(update, context)

    if step == "name":
        context.user_data["name"] = text
        user_states[uid] = "phone"
        kb = ReplyKeyboardMarkup([[KeyboardButton("📱 Отправить мой номер", request_contact=True)]],
                                 resize_keyboard=True, one_time_keyboard=True)
        await update.message.reply_text(
            f"Приятно познакомиться, {text}! 😊\n\nОтправьте телефон кнопкой ниже или "
            "напишите в формате +79991234567", reply_markup=kb)
    elif step == "phone":
        await finish_lead(update, context, text)
    else:
        await update.message.reply_text(
            "Я пока понимаю только кнопки 😅\nНажмите «📝 Оставить заявку для юриста».",
            reply_markup=keyboard_for(uid))

# ---------- Админка ----------

async def safe_edit(q, text, reply_markup=None):
    try:
        if reply_markup:
            await safe_edit(q,text=text, reply_markup=reply_markup)
        else:
            await q.edit_message_text(text=text)
    except Exception as e:
        if "Message is not modified" in str(e):
            pass
        else:
            raise e
def admin_menu(uid):
    rows = [
        [InlineKeyboardButton("📋 Лиды", callback_data="adm_leads")],
        [InlineKeyboardButton("📥 Экспорт в Excel", callback_data="adm_export")],
        [InlineKeyboardButton("🏢 Объекты", callback_data="adm_objects")],
        [InlineKeyboardButton("📚 Каталог (PDF)", callback_data="adm_catalog")],
    ]
    if is_super(uid):
        rows.append([InlineKeyboardButton("👥 Админы", callback_data="adm_admins")])
    return InlineKeyboardMarkup(rows)

async def admin_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user: return
    uid = update.effective_user.id
    if not is_admin_uid(uid):
        await update.message.reply_text("Этот раздел доступен только администраторам КомИнвест.",
                                        reply_markup=keyboard_for(uid))
        return
    await update.message.reply_text("🛠 Админ-панель КомИнвест:", )

async def publish_object(context, o):
    from html import escape as esc
    link = f"https://t.me/{BOT_USERNAME}?start=lead_obj_{o['id']}" if BOT_USERNAME else CHANNEL_URL
    caption = (f"🏢 <b>{esc(o['title'])}</b>\n\n{esc(o['description'][:650])}\n\n"
               f"💰 Цена: <b>{esc(o['price'])}</b>\n"
               f"📍 {esc(o['location'])}\n"
               f"📐 {esc(str(o.get('area', '-')))} | 🏷 {esc(str(o.get('type', '-')))}\n\n"
               f"✍️ Оставить заявку — кнопка ниже 👇")
    kb = InlineKeyboardMarkup([[InlineKeyboardButton("📝 Оставить заявку для юриста", url=link)]])
    media_ids = ([o["cover"]] if o.get("cover") else []) + o.get("photos", [])
    media_ids = media_ids[:10]
    if len(media_ids) > 1:
        media = [InputMediaPhoto(media_ids[0], caption=caption, parse_mode="HTML")] + \
                [InputMediaPhoto(p) for p in media_ids[1:]]
        await context.bot.send_media_group(CHANNEL_ID, media=media)
        await context.bot.send_message(
            CHANNEL_ID,
            text="✍️ Понравился объект? Оставьте заявку — юрист свяжется с вами 👇",
            reply_markup=kb)
    elif len(media_ids) == 1:
        await context.bot.send_photo(CHANNEL_ID, photo=media_ids[0], caption=caption,
                                     parse_mode="HTML", reply_markup=kb)
    else:
        await context.bot.send_message(CHANNEL_ID, text=caption,
                                       parse_mode="HTML", reply_markup=kb)

async def safe_edit(q, text, reply_markup=None):
    try:
        if reply_markup:
            await q.edit_message_text(text=text, reply_markup=reply_markup)
        else:
            await q.edit_message_text(text=text)
    except Exception as e:
        if "Message is not modified" in str(e):
            pass
        else:
            raise e
async def on_callback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    q = update.callback_query
    
    try:
        await q.answer()
        
        if not update.effective_user: 
            return
            
        uid = q.from_user.id
        
        # Проверка доступа к админке
        if not is_admin_uid(uid):
            await safe_edit(q, "⛔ Доступ запрещен.")
            return
            
        data = q.data
        
        # --- ЛОГИКА АДМИНКИ ---
        
        if data == "adm_enter":
            leads_count = len(load_leads())
            text = f"🔧 Админ-панель КомИнвест:\n\n• Лидов: {leads_count}\n• Объектов: {len(load_objects())}"
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton("📋 Лиды", callback_data="adm_leads"),
                 InlineKeyboardButton("🏢 Объекты", callback_data="adm_objs")],
                [InlineKeyboardButton("📊 Экспорт Excel", callback_data="adm_export"),
                 InlineKeyboardButton("👥 Админы", callback_data="adm_users")]
            ])
            await safe_edit(q, text, reply_markup=kb)

        elif data == "adm_back":
            await safe_edit(q, "🔙 Вернулись в меню.", reply_markup=admin_menu(uid))

        elif data == "adm_leads":
            leads = load_leads()
            if not leads:
                await safe_edit(q, "📋 Заявок пока нет.", reply_markup=admin_menu(uid))
                return
                
            rows = []
            for l in leads[-10:][::-1]:
                st = LEAD_STATUSES.get(l.get("status", "new"), "❓")
                btn_text = f"#{l['id']} {st} | {l['name'][:15]}..."
                rows.append([InlineKeyboardButton(btn_text, callback_data=f"lc_{l['id']}")])
            
            rows.append([InlineKeyboardButton("⬅️ В меню", callback_data="adm_back")])
            
            await safe_edit(q, f"📋 Всего заявок: {len(leads)}\nНажмите на заявку:", 
                           reply_markup=InlineKeyboardMarkup(rows))

        elif data.startswith("lc_"):
            lid = int(data.split("_")[1])
            leads = load_leads()
            lead = next((x for x in leads if x["id"] == lid), None)
            
            if not lead:
                await safe_edit(q, "Заявка не найдена.", reply_markup=admin_menu(uid))
                return
                
            card_text = (f"👤 <b>{lead['name']}</b>\n"
                         f"📞 <code>{lead['phone']}</code>\n"
                         f"📍 Источник: {lead['source']}\n"
                         f"🕒 Принят: {lead['date']}\n"
                         f"📊 Статус: <b>{LEAD_STATUSES.get(lead.get('status','new'), 'Новый')}</b>")
                         
            kb = InlineKeyboardMarkup([
                [InlineKeyboardButton("📞 В работе", callback_data=f"ls_{lid}_work"),
                 InlineKeyboardButton("✅ Закрыт", callback_data=f"ls_{lid}_done")],
                [InlineKeyboardButton("❌ Отказ", callback_data=f"ls_{lid}_fail"),
                 InlineKeyboardButton("🕐 Не дозвон", callback_data=f"ls_{lid}_recall")],
                [InlineKeyboardButton("📱 Позвонить", url=f"tel:{lead['phone']}")],
                [InlineKeyboardButton("⬅️ К списку", callback_data="adm_leads")]
            ])
            
            await safe_edit(q, card_text, reply_markup=kb, parse_mode="HTML")

        elif data.startswith("ls_"):
            parts = data.split("_")
            lid = int(parts[1])
            code = parts[2] 
            
            leads = load_leads()
            lead = next((x for x in leads if x["id"] == lid), None)
            
            if lead:
                lead["status"] = code
                lead["status_date"] = datetime.now().strftime("%Y-%m-%d %H:%M")
                update_leads(leads) 
                
                msg = f"✅ Заявка #{lid}: {LEAD_STATUSES[code]}"
                await safe_edit(q, msg, reply_markup=InlineKeyboardMarkup([
                    [InlineKeyboardButton("Вернуться к карточке", callback_data=f"lc_{lid}")],
                    [InlineKeyboardButton("К списку", callback_data="adm_leads")]
                ]))
            else:
                await safe_edit(q, "Ошибка: заявка не найдена.", reply_markup=admin_menu(uid))

        elif data == "adm_objs":
            objs = load_objects()
            rows = [[InlineKeyboardButton(f"#{o['id']} {o['title'][:20]}...", callback_data=f"ob_{o['id']}")] for o in objs[-5:]]
            rows.append([InlineKeyboardButton("➕ Добавить объект", callback_data="obj_add")])
            rows.append([InlineKeyboardButton("⬅️ В меню", callback_data="adm_back")])
            
            await safe_edit(q, f"🏢 Объектов: {len(objs)}", reply_markup=InlineKeyboardMarkup(rows))

        elif data == "adm_export":
            try:
                wb = Workbook()
                ws = wb.active
                ws.title = "Leads"
                ws.append(["№", "Имя", "Телефон", "Источник", "Дата", "Статус", "Изменён"])
                
                for l in load_leads():
                    st = LEAD_STATUSES.get(l.get("status", "new"), l.get("status", "new"))
                    ws.append([l["id"], l["name"], l["phone"], l["source"], l["date"], st, l.get("status_date", "")])
                    
                buf = io.BytesIO()
                wb.save(buf)
                buf.seek(0)
                
                await context.bot.send_document(
                    chat_id=q.from_user.id,
                    document=buf,
                    filename=f"kominvest_leads_{datetime.now().strftime('%d%m%Y_%H%M')}.xlsx",
                    caption="📊 Выгрузка лидов со статусами."
                )
                await q.answer("Файл отправлен!", show_alert=True)
            except Exception as e:
                print(f"Export error: {e}")
                await q.answer("Ошибка при создании файла.", show_alert=True)

        else:
            await safe_edit(q, "🔙 Действие не распознано. Вернулись в меню.", reply_markup=admin_menu(uid))

    except telegram.error.BadRequest as e:
        err_str = str(e).lower()
        if "message is not modified" in err_str:
            pass 
        elif "bot was blocked by the user" in err_str:
            pass 
        else:
            print(f"Telegram API Error in callback: {e}")
            
    except Exception as e:
        print(f"Critical error in on_callback: {e}")
        import traceback
        traceback.print_exc()
        try:
            await q.answer("Произошла внутренняя ошибка.", show_alert=True)
        except:
            pass


async def leads_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user or not is_admin_uid(update.effective_user.id): return
    leads = load_leads()
    if not leads:
        await update.message.reply_text("📋 Заявок пока нет.")
        return
    lines = [f"📋 Заявки: всего {len(leads)}\n"]
    for l in leads[-15:]:
        lines.append(f"#{l['id']} | {l['name']} | {l['phone']} | {l['source']} | {l['date']}")
    await update.message.reply_text("\n".join(lines))

async def stats_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user or not is_admin_uid(update.effective_user.id): return
    leads = load_leads()
    today = datetime.now().strftime("%Y-%m-%d")
    objs = load_objects()
    await update.message.reply_text(
        f"📊 Статистика КомИнвест:\n\nВсего заявок: {len(leads)}\n"
        f"Сегодня: {sum(1 for l in leads if l['date'].startswith(today))}\n"
        f"Объектов: {len(objs)} | Опубликовано: {sum(1 for o in objs if o.get('published'))}")

async def export_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user or not is_admin_uid(update.effective_user.id): return
    import openpyxl
    from openpyxl.styles import Font
    leads = load_leads()
    if not leads:
        await update.message.reply_text("Экспортировать нечего.")
        return
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "Заявки"
    ws.append(["№", "Имя", "Телефон", "Источник", "Дата", "Статус", "Статус изменён"])
    for cell in ws[1]: cell.font = Font(bold=True)
    for l in leads:
        ws.append([l["id"], l["name"], l["phone"], l["source"], l["date"], LEAD_STATUSES.get(l.get("status", "new"), l.get("status", "new")), l.get("status_date", "")])
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    await update.message.reply_document(document=buf,
        filename=f"kominvest_leads_{datetime.now().strftime('%d%m%Y')}.xlsx",
        caption=f"📥 Выгрузка заявок: {len(leads)} шт.")

PHOTO_CACHE = os.path.join("data", "photos")

async def get_photo_bytes(bot, file_id):
    if not file_id: return None
    os.makedirs(PHOTO_CACHE, exist_ok=True)
    path = os.path.join(PHOTO_CACHE, file_id.replace("/", "_") + ".jpg")
    if os.path.exists(path):
        with open(path, "rb") as f: return f.read()
    try:
        f = await bot.get_file(file_id)
        data = bytes(await f.download_as_bytearray())
        with open(path, "wb") as fp: fp.write(data)
        return data
    except Exception as e:
        print(f"photo download error: {e}")
        return None

async def build_catalog_bytes(context):
    objs = [dict(o) for o in load_objects() if o.get("category") == "commercial"]
    if not objs: return None
    for o in objs:
        o["_cover_bytes"] = await get_photo_bytes(context.bot, o.get("cover"))
    return catalog.build_catalog(objs)

async def catalog_cmd(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not update.effective_user or not is_admin_uid(update.effective_user.id): return
    m = await update.message.reply_text("⏳ Собираю каталог...")
    pdf = await build_catalog_bytes(context)
    if not pdf:
        await m.edit_text("Нет объектов для каталога.")
        return
    await update.message.reply_document(io.BytesIO(pdf),
        filename="kominvest_catalog.pdf",
        caption="📚 Каталог объектов КомИнвест")
    try: await m.delete()
    except Exception: pass

async def error_handler(update: Update, context: ContextTypes.DEFAULT_TYPE):
    traceback.print_exc()
    try:
        if update is not None and update.effective_message is not None:
            await update.effective_message.reply_text(f"⚠️ Техническая ошибка: {context.error}")
    except Exception:
        pass

# ---------- Запуск ----------

if __name__ == "__main__":
    import asyncio
    from aiohttp import web

    app = Application.builder().token(BOT_TOKEN).build()
    app.add_handler(CommandHandler("start", start))
    app.add_handler(CommandHandler("admin", admin_cmd))
    app.add_handler(CommandHandler("leads", leads_cmd))
    app.add_handler(CommandHandler("stats", stats_cmd))
    app.add_handler(CommandHandler("export", export_cmd))
    app.add_handler(CallbackQueryHandler(on_callback))
    app.add_handler(MessageHandler(filters.CONTACT, handle_contact))
    app.add_handler(MessageHandler(filters.PHOTO, handle_photo))
    app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
    app.add_handler(CommandHandler("catalog", catalog_cmd))
    app.add_error_handler(error_handler)

    if WEBHOOK_URL:
        async def webhook_handler(request):
            try:
                d = await request.json()
            except Exception:
                return web.Response(status=400, text="bad json")
            upd = Update.de_json(d, app.bot)
            if upd is not None:
                await app.update_queue.put(upd)
            return web.Response(text="OK")

        async def health_handler(request):
            return web.Response(text="OK")

        async def main():
            global BOT_USERNAME
            gh_store.sync_all({"leads.json": [], "objects.json": [],
                               "admins.json": {"super": [ADMIN_ID], "admins": []}})
            await app.initialize()
            await app.start()
            BOT_USERNAME = app.bot.username
            print(f"CHANNEL_ID={CHANNEL_ID} | BOT=@{BOT_USERNAME}")
            try:
                await app.bot.set_webhook(WEBHOOK_URL, allowed_updates=Update.ALL_TYPES)
                print(f"Webhook зарегистрирован: {WEBHOOK_URL}")
            except Exception as e:
                print(f"Не удалось зарегистрировать webhook: {e}")
            web_app = web.Application()
            web_app.router.add_post("/", webhook_handler)
            web_app.router.add_get("/", health_handler)
            runner = web.AppRunner(web_app)
            await runner.setup()
            site = web.TCPSite(runner, "0.0.0.0", PORT)
            await site.start()
            print(f"Webhook mode. URL: {WEBHOOK_URL} (GET / = 200 OK)")
            await asyncio.Event().wait()

        print("Запускаю webhook-сервер...")
        asyncio.run(main())
    else:
        print("Polling mode (локально)...")
        app.run_polling(allowed_updates=Update.ALL_TYPES)